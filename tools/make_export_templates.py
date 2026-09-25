"""Build the styled export templates in matlab/templates/.

Dev-only (needs openpyxl); the app never runs this. The committed .xlsx
files are the output, and MATLAB fills them with values at export time
(writecell, UseExcel=false, PreserveFormat=true, AutoFitWidth=false).

Every block MATLAB writes into is a named range. The MATLAB side reads the
same names from report.exportLayout, and tExportTemplates fails if the two
ever disagree.

    python tools/make_export_templates.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

OUT = Path(__file__).resolve().parents[1] / "matlab" / "templates"

# Same values as gui.palette, so the workbook and the app agree.
PASS_BG, PASS_FG = "C7F0C7", "006600"
FAIL_BG, FAIL_FG = "FFC7C7", "CC0000"
NOTEVAL_BG, NOTEVAL_FG = "FFF3CD", "856404"
HEADER_BG = "E0ECF9"
MUTED = "666666"
RULE = "C7C7CC"

MARGIN_FMT = "+0.00;-0.00;0.00"
thin = Side(style="thin", color=RULE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)


def status_rules(ws, rng):
    for text, bg, fg in (("FAIL", FAIL_BG, FAIL_FG),
                         ("Pass", PASS_BG, PASS_FG),
                         ("Not evaluated", NOTEVAL_BG, NOTEVAL_FG)):
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=[f'"{text}"'],
            fill=fill(bg), font=Font(color=fg, bold=True)))


def header(ws, cells):
    for ref, text in cells:
        c = ws[ref]
        c.value = text
        c.font = Font(bold=True)
        c.fill = fill(HEADER_BG)
        c.border = box
        c.alignment = Alignment(vertical="center")


def grid(ws, cols, rows, wrap=False):
    for r in rows:
        for col in cols:
            c = ws[f"{col}{r}"]
            c.border = box
            c.alignment = Alignment(vertical="top" if wrap else "center",
                                    wrap_text=wrap)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def name(wb, key, ref):
    wb.defined_names[key] = DefinedName(key, attr_text=ref)


def single():
    wb = Workbook()

    # ---- Slide: one block to paste onto a 16:9 slide --------------------
    ws = wb.active
    ws.title = "Slide"
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 20, "B": 34, "C": 2, "D": 24, "E": 13, "F": 14,
                "G": 34, "I": 10})
    ws.column_dimensions["I"].hidden = True

    ws.merge_cells("A1:G1")
    ws["A1"].font = Font(size=16, bold=True)
    ws.row_dimensions[1].height = 24
    for cls, color in (("fail", FAIL_FG), ("noteval", NOTEVAL_FG),
                       ("pass", PASS_FG)):
        ws.conditional_formatting.add("A1", FormulaRule(
            formula=[f'$I$1="{cls}"'], font=Font(color=color, bold=True)))

    ws.merge_cells("A2:G2")
    ws["A2"].font = Font(italic=True, color=MUTED)

    header(ws, [("A4", "Input"), ("B4", "Value"), ("D4", "Check"),
                ("E4", "MS / R"), ("F4", "Status"),
                ("G4", "Governing equation")])
    grid(ws, "AB", range(5, 20))
    grid(ws, "DEFG", range(5, 20))
    for r in range(5, 20):
        ws[f"A{r}"].font = Font(bold=True)
        ws[f"E{r}"].number_format = MARGIN_FMT
        ws[f"E{r}"].alignment = Alignment(horizontal="right",
                                          vertical="center")
    status_rules(ws, "F5:F19")

    ws.merge_cells("A21:G21")
    ws["A21"].font = Font(italic=True, size=9, color=MUTED)
    ws["A21"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[21].height = 36

    name(wb, "SlideTitle", "Slide!$A$1")
    name(wb, "SlideSubtitle", "Slide!$A$2")
    name(wb, "SlideVerdictClass", "Slide!$I$1")
    name(wb, "SlideInputs", "Slide!$A$5:$B$19")
    name(wb, "SlideMargins", "Slide!$D$5:$G$19")
    name(wb, "SlideScope", "Slide!$A$21")

    # ---- Detail: the checker's view of every check ----------------------
    ws = wb.create_sheet("Detail")
    ws.freeze_panes = "A2"
    widths(ws, {"A": 24, "B": 14, "C": 16, "D": 60, "E": 48, "F": 48})
    header(ws, [("A1", "Check"), ("B1", "Status"), ("C1", "MS / R"),
                ("D1", "Governing equation"),
                ("E1", "Inputs (substituted)"), ("F1", "Detail")])
    grid(ws, "ABCDEF", range(2, 17), wrap=True)
    for r in range(2, 17):
        ws[f"C{r}"].number_format = MARGIN_FMT
    status_rules(ws, "B2:B16")
    name(wb, "DetailRows", "Detail!$A$2:$F$16")

    # ---- About: provenance ------------------------------------------------
    ws = wb.create_sheet("About")
    widths(ws, {"A": 22, "B": 100})
    header(ws, [("A1", "Item"), ("B1", "Value")])
    grid(ws, "AB", range(2, 26), wrap=True)
    for r in range(2, 26):
        ws[f"A{r}"].font = Font(bold=True)
    name(wb, "AboutRows", "About!$A$2:$B$25")

    wb.save(OUT / "export_single.xlsx")


if __name__ == "__main__":
    single()
    print("wrote", OUT / "export_single.xlsx")
